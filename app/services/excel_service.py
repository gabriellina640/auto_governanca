from __future__ import annotations

import getpass
import logging
import os
import shutil
import tempfile
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

DEFAULT_SHEET_NAME = "Banco de dados"


class ExcelServiceError(Exception):
    pass


class ExcelService:
    def __init__(self, config: dict[str, Any]) -> None:
        self.config = config
        self.path: Path | None = None
        self.workbook = None
        self.sheet: Worksheet | None = None
        self.headers: list[str] = []
        self.loaded_mtime: float | None = None

    def load(self, path: str | Path) -> None:
        file_path = Path(path).expanduser()
        if not file_path.exists():
            raise ExcelServiceError("Arquivo não encontrado.")
        if file_path.suffix.lower() != ".xlsx":
            raise ExcelServiceError("Use um arquivo .xlsx. Arquivos .xls não são suportados neste MVP.")

        try:
            self.workbook = load_workbook(file_path)
        except PermissionError as exc:
            raise ExcelServiceError("Não foi possível abrir. Feche o Excel e tente novamente.") from exc
        except Exception as exc:
            raise ExcelServiceError(f"Erro ao carregar Excel: {exc}") from exc

        self.sheet = self._select_sheet()

        self.path = file_path
        self.loaded_mtime = file_path.stat().st_mtime
        self.headers = self._read_headers()
        if not self.headers:
            raise ExcelServiceError("A primeira linha da planilha precisa conter os cabeçalhos.")

    def _read_headers(self) -> list[str]:
        if not self.sheet:
            return []
        headers: list[str] = []
        for cell in self.sheet[1]:
            value = "" if cell.value is None else str(cell.value).strip()
            headers.append(value)
        while headers and headers[-1] == "":
            headers.pop()
        if any(header == "" for header in headers):
            raise ExcelServiceError("A primeira linha não pode ter cabeçalhos vazios entre colunas preenchidas.")
        if len(headers) != len(set(headers)):
            raise ExcelServiceError("A primeira linha não pode ter cabeçalhos duplicados.")
        return headers

    def _select_sheet(self) -> Worksheet:
        if not self.workbook:
            raise ExcelServiceError("Nenhum arquivo carregado.")

        sheet_name = self.config.get("sheet_name") or DEFAULT_SHEET_NAME
        if sheet_name in self.workbook.sheetnames:
            return self.workbook[sheet_name]
        if DEFAULT_SHEET_NAME in self.workbook.sheetnames:
            return self.workbook[DEFAULT_SHEET_NAME]
        return self.workbook[self.workbook.sheetnames[0]]

    def get_rows(self) -> list[list[Any]]:
        if not self.sheet:
            return []
        rows: list[list[Any]] = []
        max_col = len(self.headers)
        for row in self.sheet.iter_rows(min_row=2, max_col=max_col, values_only=True):
            if any(value is not None and str(value).strip() != "" for value in row):
                rows.append(list(row))
        return rows

    def _configured_headers(self) -> list[str]:
        configured = self.config.get("form_fields") or []
        if not configured:
            return []
        return [header for header in configured if header in self.headers]

    def get_visible_headers(self) -> list[str]:
        configured = self._configured_headers()
        return configured or self.headers

    def get_visible_rows(self) -> list[list[Any]]:
        rows = self.get_rows()
        visible_headers = self.get_visible_headers()
        if visible_headers == self.headers:
            return [
                [self._format_display_value(header, value) for header, value in zip(self.headers, row, strict=False)]
                for row in rows
            ]

        indexes = [self.headers.index(header) for header in visible_headers]
        return [
            [
                self._format_display_value(header, row[index] if index < len(row) else "")
                for header, index in zip(visible_headers, indexes, strict=False)
            ]
            for row in rows
        ]

    def get_form_headers(self) -> list[str]:
        hidden = set(self.config.get("hidden_form_fields", []))
        auto_cols = set((self.config.get("auto_columns") or {}).values())
        base_headers = self._configured_headers() or self.headers
        return [h for h in base_headers if h not in hidden and h not in auto_cols]

    def has_external_change(self) -> bool:
        if not self.path or self.loaded_mtime is None or not self.path.exists():
            return False
        return self.path.stat().st_mtime > self.loaded_mtime + 0.5

    def add_record(self, values: dict[str, Any]) -> None:
        if not self.sheet:
            raise ExcelServiceError("Nenhuma planilha carregada.")

        auto = self.config.get("auto_columns") or {}
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        user = getpass.getuser()

        record = {header: values.get(header, "") for header in self.headers}

        id_col = auto.get("id")
        timestamp_col = auto.get("timestamp")
        user_col = auto.get("user")

        if id_col in self.headers:
            record[id_col] = self._next_id(id_col)
        if timestamp_col in self.headers:
            record[timestamp_col] = now
        if user_col in self.headers:
            record[user_col] = user

        template_row = self._last_data_row()
        self.sheet.append([self._coerce_value(header, record.get(header, "")) for header in self.headers])
        new_row = self.sheet.max_row
        self._copy_row_format(template_row, new_row)
        self._apply_row_formats(new_row)
        self._expand_tables(new_row)
        logging.info("Registro incluído por %s", user)

    def _date_fields(self) -> set[str]:
        return set(self.config.get("date_fields", []))

    def _format_display_value(self, header: str, value: Any) -> Any:
        if header in self._date_fields() and isinstance(value, datetime | date):
            return value.strftime(self.config.get("display_date_format", "%d/%m/%Y"))
        return value

    def _coerce_value(self, header: str, value: Any) -> Any:
        if header not in self._date_fields() or not isinstance(value, str):
            return value

        stripped = value.strip()
        if not stripped:
            return ""

        try:
            return datetime.strptime(stripped, "%d/%m/%Y")
        except ValueError:
            return value

    def _apply_row_formats(self, row_index: int) -> None:
        if not self.sheet:
            return

        date_fields = self._date_fields()
        excel_date_format = self.config.get("excel_date_format", "DD/MM/YYYY")
        for column_index, header in enumerate(self.headers, start=1):
            if header in date_fields:
                self.sheet.cell(row=row_index, column=column_index).number_format = excel_date_format

    def _last_data_row(self) -> int:
        if not self.sheet:
            return 1

        max_col = len(self.headers)
        for row_index in range(self.sheet.max_row, 1, -1):
            row_has_data = any(
                self.sheet.cell(row=row_index, column=column_index).value not in (None, "")
                for column_index in range(1, max_col + 1)
            )
            if row_has_data:
                return row_index
        return 1

    def _copy_row_format(self, source_row: int, target_row: int) -> None:
        if not self.sheet or source_row == target_row:
            return

        self.sheet.row_dimensions[target_row].height = self.sheet.row_dimensions[source_row].height

        for column_index in range(1, len(self.headers) + 1):
            source_cell = self.sheet.cell(row=source_row, column=column_index)
            target_cell = self.sheet.cell(row=target_row, column=column_index)
            if source_cell.has_style:
                target_cell._style = copy(source_cell._style)
            if source_cell.number_format:
                target_cell.number_format = source_cell.number_format
            if source_cell.alignment:
                target_cell.alignment = copy(source_cell.alignment)
            if source_cell.border:
                target_cell.border = copy(source_cell.border)
            if source_cell.fill:
                target_cell.fill = copy(source_cell.fill)
            if source_cell.font:
                target_cell.font = copy(source_cell.font)
            if source_cell.protection:
                target_cell.protection = copy(source_cell.protection)

    def _expand_tables(self, new_row: int) -> None:
        if not self.sheet:
            return

        for table in self.sheet.tables.values():
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            if min_row <= 1 and new_row == max_row + 1 and min_col == 1:
                table.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{new_row}"

    def _next_id(self, id_header: str) -> int:
        if not self.sheet:
            return 1
        col_index = self.headers.index(id_header) + 1
        max_id = 0
        for row in range(2, self.sheet.max_row + 1):
            value = self.sheet.cell(row=row, column=col_index).value
            try:
                max_id = max(max_id, int(value))
            except (TypeError, ValueError):
                continue
        return max_id + 1

    def backup(self) -> Path | None:
        if not self.path or not self.path.exists():
            return None
        if not self.config.get("backup_enabled", True):
            return None

        backup_folder_name = self.config.get("backup_folder", "backups")
        backup_dir = self.path.parent / backup_folder_name
        backup_dir.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_path = backup_dir / f"{self.path.stem}_backup_{timestamp}{self.path.suffix}"
        shutil.copy2(self.path, backup_path)
        logging.info("Backup criado: %s", backup_path)
        return backup_path

    def save(self) -> None:
        if not self.path or not self.workbook:
            raise ExcelServiceError("Nenhum arquivo carregado.")

        self.backup()

        fd, temp_name = tempfile.mkstemp(
            prefix=f".{self.path.stem}_",
            suffix=".xlsx",
            dir=self.path.parent,
        )
        os.close(fd)
        temp_path = Path(temp_name)
        try:
            self.workbook.save(temp_path)
            os.replace(temp_path, self.path)
            self.loaded_mtime = self.path.stat().st_mtime
            logging.info("Arquivo salvo: %s", self.path)
        except PermissionError as exc:
            raise ExcelServiceError("Não consegui salvar. Feche o arquivo no Excel e tente novamente.") from exc
        except Exception as exc:
            raise ExcelServiceError(f"Erro ao salvar Excel: {exc}") from exc
        finally:
            if temp_path.exists():
                try:
                    temp_path.unlink()
                except Exception:
                    pass
